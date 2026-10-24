"""Tidy Census 2022 education / migration tables relevant to student mobility.

Inputs (data/external/education/census2022_education_migration/):
  ED18.xlsx  currently attending educational institution by level x island (usual residence)
  ED17.xlsx  population / currently studying by age group x island
  MG13.xlsx  usual residents of Male' by island of registration
  Migration-Report-Census-2022.pdf  tables 5.6 (p61), 5.8 (p66), 5.9 (p67), 5.10 (p68)
Outputs: tidy CSVs in the same folder.

Usage: python -I scripts/edu_tidy_census2022.py
"""
import csv
import re
import sys
from pathlib import Path

import openpyxl
import pymupdf

sys.stdout.reconfigure(encoding="utf-8")
B = Path(r"C:\Users\user\Music\Datathon\data\external\education\census2022_education_migration")
URL = "https://statisticsmaldives.gov.mv/mbs/wp-content/uploads/2023/09/{}.xlsx"
AGG = {"Republic", "Maale", "Atolls (Admin and Non- Admin Islands)", "Admin islands",
       "Non-Admin islands", "Non- Admin islands", "Total"}


def write(path, rows):
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0]))
        w.writeheader()
        w.writerows(rows)
    print(f"{path.name}: {len(rows)} rows")


def num(v):
    if v is None or v == "" or v == "-":
        return 0
    try:
        return float(v) if isinstance(v, float) and not float(v).is_integer() else int(float(v))
    except (TypeError, ValueError):
        return None


def ed18():
    ws = openpyxl.load_workbook(B / "ED18.xlsx", data_only=True).worksheets[0]
    cols = ["total_attending", "pre_school", "primary_g1_7", "lower_secondary_g8_10",
            "higher_secondary_g11_12", "o_level_classes", "a_level_classes", "foundation",
            "diploma", "bachelors", "masters_plus", "phd", "certificate_sanadhu", "special_needs"]
    out = []
    started = False
    for r in ws.iter_rows(values_only=True):
        if r[0] == "(1)":
            started = True
            continue
        if not started or r[0] is None:
            continue
        a = str(r[0]).strip()
        level = "aggregate" if r[1] is None else "island"
        atoll, island, vals = a, ("" if r[1] is None else str(r[1]).strip()), r[2:16]
        if not isinstance(vals[0], (int, float)):
            continue
        rec = dict(atoll=atoll, island=island, row_type=level)
        rec.update({c: num(v) for c, v in zip(cols, vals)})
        rec["tertiary_attending"] = sum(rec[c] or 0 for c in ("foundation", "diploma", "bachelors", "masters_plus", "phd"))
        rec["source_url"] = URL.format("ED18")
        out.append(rec)
    write(B / "census2022_ED18_attending_by_level_island.csv", out)


def ed17():
    ws = openpyxl.load_workbook(B / "ED17.xlsx", data_only=True).worksheets[0]
    ages = ["2-5", "6-12", "13-15", "16-19", "20-24", "25-29", "30-34", "35-39", "40-44", "45-49", "50+"]
    out = []
    started = False
    for r in ws.iter_rows(values_only=True):
        if r[0] == "(1)":
            started = True
            continue
        if not started or r[0] is None:
            continue
        a = str(r[0]).strip()
        atoll, island, vals = a, ("" if r[1] is None else str(r[1]).strip()), list(r[2:35])
        if not isinstance(vals[0], (int, float)):
            continue
        for k, ag in enumerate(ages):
            pop, stu = num(vals[k]), num(vals[11 + k])
            out.append(dict(atoll=atoll, island=island, age_group=ag, resident_maldivian_pop=pop,
                            currently_studying=stu,
                            pct_studying=round(100 * stu / pop, 1) if pop else None,
                            source_url=URL.format("ED17")))
    write(B / "census2022_ED17_studying_by_age_island_long.csv", out)


def mg13():
    ws = openpyxl.load_workbook(B / "MG13.xlsx", data_only=True).worksheets[0]
    out = []
    started = False
    for r in ws.iter_rows(values_only=True):
        if r[0] == "(1)":
            started = True
            continue
        if not started or r[0] is None:
            continue
        a = str(r[0]).strip()
        isl, vals = ("" if r[1] is None else str(r[1]).strip()), r[2:5]
        if not isinstance(vals[0], (int, float)):
            continue
        out.append(dict(registered_atoll=a, registered_island=isl, male_residents_both=num(vals[0]),
                        male_residents_male=num(vals[1]), male_residents_female=num(vals[2]),
                        source_url="https://statisticsmaldives.gov.mv/mbs/wp-content/uploads/2023/08/MG13.xlsx"))
    write(B / "census2022_MG13_male_residents_by_registered_island.csv", out)


def pct(s):
    m = re.match(r"(\d+)%", (s or "").strip())
    return int(m.group(1)) if m else None


def report_tables():
    pdf = B / "Migration-Report-Census-2022.pdf"
    d = pymupdf.open(pdf)
    src = "https://statisticsmaldives.gov.mv/mbs/wp-content/uploads/2024/02/Migration-Report-Census-2022.pdf"
    reasons = ["for_education", "new_job", "family_migration", "marriage", "parent_guardian", "medical"]
    out = []
    for pno, period in ((66, "all migrants (lifetime, to Maale)"), (67, "recent migrants (<1 year, to Maale)")):
        rows = d[pno - 1].find_tables().tables[0].extract()
        for r in rows[1:]:
            rec = dict(origin_atoll=r[0], destination="Maale", migrant_group=period)
            rec.update({k: pct(v) for k, v in zip(reasons, r[1:])})
            rec["source"] = f"MBS Migration Report Census 2022, Table 5.{8 if pno == 66 else 9}, p{pno}"
            rec["source_url"] = src
            out.append(rec)
    write(B / "census2022_reason_for_migration_to_male_by_atoll_pct.csv", out)

    out = []
    rows = d[60].find_tables().tables[0].extract()
    labels = ["for_education", "new_job", "return_after_education", "return_to_live",
              "family_migration", "marriage", "parent_guardian", "medical"]
    heads = [("<1 year", "Republic"), ("<1 year", "Maale"), ("<1 year", "Atolls"), ("<1 year", "Admin islands"),
             ("7 years ago", "Republic"), ("7 years ago", "Maale"), ("7 years ago", "Atolls"), ("7 years ago", "Admin islands")]
    for lab, r in zip(labels, rows[2:]):
        for (period, loc), v in zip(heads, r[1:]):
            out.append(dict(reason=lab, period=period, destination_locality=loc, pct=pct(v),
                            source="MBS Migration Report Census 2022, Table 5.6, p61", source_url=src))
    rows = d[67].find_tables().tables[0].extract()
    cities = [("<1 year", "Kulhudhuffushi"), ("7 years ago", "Kulhudhuffushi"), ("<1 year", "Fuvahmulah City"),
              ("7 years ago", "Fuvahmulah City"), ("<1 year", "Addu City"), ("7 years ago", "Addu City")]
    for lab, r in zip(labels, rows[2:]):
        for (period, loc), v in zip(cities, r[1:]):
            out.append(dict(reason=lab, period=period, destination_locality=loc, pct=pct(v),
                            source="MBS Migration Report Census 2022, Table 5.10, p68", source_url=src))
    write(B / "census2022_reason_for_migration_by_destination_pct.csv", out)


if __name__ == "__main__":
    ed18()
    ed17()
    mg13()
    report_tables()
