"""Print the non-empty cells of each sheet of one or more xlsx files (inspection helper).

Usage: python -I scripts/edu_xlsx_peek.py <xlsx> [<xlsx> ...] [--rows N]
"""
import sys
sys.stdout.reconfigure(encoding="utf-8")

import openpyxl

args = sys.argv[1:]
nrows = 60
if "--rows" in args:
    i = args.index("--rows")
    nrows = int(args[i + 1])
    args = args[:i] + args[i + 2:]
for p in args:
    wb = openpyxl.load_workbook(p, data_only=True, read_only=True)
    for ws in wb.worksheets:
        print(f"##### {p} :: {ws.title} ({ws.max_row}x{ws.max_column})")
        k = 0
        for row in ws.iter_rows(values_only=True):
            vals = [str(v).strip() for v in row if v is not None and str(v).strip()]
            if vals:
                print(" | ".join(vals)[:300])
                k += 1
                if k >= nrows:
                    print("...")
                    break
