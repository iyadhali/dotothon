"""Tidy selected MBS Statistical Yearbook education tables (downloaded XLSX) into CSVs.

Inputs  (data/external/education/mbs_statistical_yearbook/):
  yb2024/7.2.xlsx  Student enrolment by SCHOOL x grade x sex, 2022-2023 (with atoll & island)
  yb2023/7.7.xlsx  Student enrolment by ISLAND x grade x sex, 2021-2022 (as of 1 Jun 2022)
  yb2022/7.29.xlsx Student enrolment in universities & colleges by institution, 2019-2021
Outputs written next to inputs.

Usage: python -I scripts/edu_tidy_mbs_yearbook.py
"""
import csv
import re
import sys
from pathlib import Path

import openpyxl

sys.stdout.reconfigure(encoding="utf-8")
B = Path(r"C:\Users\user\Music\Datathon\data\external\education\mbs_statistical_yearbook")
GR = ["BN", "N", "LKG", "UKG"] + [f"GR{i}" for i in range(1, 13)]
ATOLL_CODES = {
    "Male city": "K-Male", "HA": "HA", "HDh": "HDh", "Sh": "Sh", "N": "N", "R": "R", "B": "B",
    "Lh": "Lh", "K": "K", "AA": "AA", "ADh": "ADh", "V": "V", "M": "M", "F": "F", "Dh": "Dh",
    "Th": "Th", "L": "L", "GA": "GA", "GDh": "GDh", "Gn": "Gn", "Sn": "S", "S": "S", "Adh": "ADh",
}


def iv(v):
    if v is None or (isinstance(v, str) and not v.strip()):
        return 0
    return int(round(float(v)))


def derive(rec):
    present = [g for g in GR if rec[f"{g}_M"] + rec[f"{g}_F"] > 0]
    rec["lowest_grade"] = present[0] if present else ""
    rec["highest_grade"] = present[-1] if present else ""
    rec["grade_11_12"] = sum(rec[f"GR{i}_{s}"] for i in (11, 12) for s in "MF")
    rec["offers_higher_secondary"] = rec["grade_11_12"] > 0
    rec["grades_total"] = sum(rec[f"{g}_{s}"] for g in GR for s in "MF")


def school_2022_23():
    ws = openpyxl.load_workbook(B / "yb2024" / "7.2.xlsx", data_only=True).worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    out = []
    for r in rows[5:]:
        if not isinstance(r[0], (int, float)):
            continue
        rec = dict(year="2022-2023", sr=int(r[0]), atoll_label=r[1], atoll_code=ATOLL_CODES.get(r[1], r[1]),
                   school=str(r[2]).strip(), island=str(r[3]).strip(),
                   total=iv(r[4]), total_F=iv(r[5]), total_M=iv(r[6]))
        for k, g in enumerate(GR):
            rec[f"{g}_F"] = iv(r[7 + 2 * k])
            rec[f"{g}_M"] = iv(r[8 + 2 * k])
        derive(rec)
        rec["check_ok"] = rec["grades_total"] == rec["total"]
        rec["source_url"] = ("https://statisticsmaldives.gov.mv/yearbook/2024/wp-content/uploads/"
                             "sites/11/2024/11/7.2.xlsx")
        out.append(rec)
    f = B / "yb2024" / "school_enrolment_by_grade_2022_23.csv"
    write(f, out)
    bad = sum(not r["check_ok"] for r in out)
    print(f"{f.name}: {len(out)} schools, {sum(r['total'] for r in out)} students, "
          f"{sum(r['offers_higher_secondary'] for r in out)} offer Gr11-12, sumcheck fails={bad}")


def island_2021_22():
    ws = openpyxl.load_workbook(B / "yb2023" / "7.7.xlsx", data_only=True).worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    out = []
    atoll = None
    hdr_seen = False
    for r in rows:
        if any(isinstance(v, str) and v.strip() == "Both Sexes" for v in r):
            hdr_seen = True
            continue
        name = r[0]
        if name is None:
            continue
        name = str(name).strip()
        if not hdr_seen or not isinstance(r[1], (int, float)):
            continue
        if name.lower().startswith("total"):
            continue
        # Atoll header rows are short codes (HA, HDh...) or "Greater Male' Area"
        is_atoll = name in ATOLL_CODES or name.startswith("Greater") or name in ("Sn", "S")
        if is_atoll:
            atoll = name
        rec = dict(year="2021-2022", atoll=atoll, island=name, is_atoll_total=is_atoll,
                   total=iv(r[1]), total_M=iv(r[2]), total_F=iv(r[3]))
        for k, g in enumerate(GR):
            rec[f"{g}_M"] = iv(r[4 + 2 * k])
            rec[f"{g}_F"] = iv(r[5 + 2 * k])
        derive(rec)
        rec["check_ok"] = rec["grades_total"] == rec["total"]
        rec["source_url"] = ("https://statisticsmaldives.gov.mv/yearbook/2023/wp-content/uploads/"
                             "sites/10/2024/05/7.7.xlsx")
        out.append(rec)
    f = B / "yb2023" / "island_enrolment_by_grade_2021_22.csv"
    write(f, out)
    isl = [r for r in out if not r["is_atoll_total"]]
    print(f"{f.name}: {len(isl)} islands, atoll rows={len(out)-len(isl)}, "
          f"fails={sum(not r['check_ok'] for r in out)}")


def hei_2019_2021():
    ws = openpyxl.load_workbook(B / "yb2022" / "7.29.xlsx", data_only=True).worksheets[0]
    out = []
    year = None
    for r in ws.iter_rows(values_only=True):
        vals = [v for v in r if v is not None]
        if not vals:
            continue
        s = str(vals[0])
        m = re.search(r"COLLEGES,\s*(\d{4})", s)
        if m:
            year = int(m.group(1))
            continue
        if len(vals) >= 3 and isinstance(vals[1], (int, float)) and s.strip() != "Total":
            out.append(dict(year=year, institution=s.strip(), enrolment=int(round(vals[1])),
                            pct_of_total=round(float(vals[2]), 2),
                            source="Ministry of Higher Education via MBS Statistical Yearbook 2022 Table 7.29",
                            source_url="https://statisticsmaldives.gov.mv/yearbook/2022/wp-content/uploads/sites/9/2023/03/7.29.xlsx"))
    f = B / "yb2022" / "hei_enrolment_by_institution_2019_2021.csv"
    write(f, out)
    print(f"{f.name}: {len(out)} rows")


def write(path, rows):
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0]))
        w.writeheader()
        w.writerows(rows)


if __name__ == "__main__":
    school_2022_23()
    island_2021_22()
    hei_2019_2021()
