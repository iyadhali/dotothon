"""Profile MTCC.xlsx (all sheets). Usage: python -I profile_mtcc.py <path-to-xlsx>"""
import sys
import openpyxl
import pandas as pd

pd.set_option("display.width", 250)
pd.set_option("display.max_columns", 50)
pd.set_option("display.max_colwidth", 60)

path = sys.argv[1]
wb = openpyxl.load_workbook(path, read_only=False, data_only=True)
for ws in wb.worksheets:
    print("=" * 100)
    print(f"SHEET: {ws.title!r} dims={ws.dimensions} max_row={ws.max_row} max_col={ws.max_column}")
    print("merged:", list(ws.merged_cells.ranges)[:20])
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i >= (int(sys.argv[2]) if len(sys.argv) > 2 else 400):
            print("... truncated")
            break
        if any(v is not None for v in row):
            print(i + 1, "|", " | ".join("" if v is None else str(v) for v in row))
